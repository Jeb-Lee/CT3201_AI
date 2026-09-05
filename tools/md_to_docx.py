"""Markdown -> .docx for the CT3201 report drafts.

    pip install python-docx
    python tools/md_to_docx.py

Rebuilds docs/report/CT3201_Report.docx and CT3201_AI_Interaction_Log.docx
from the matching .md files.
"""
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
NOTE = RGBColor(0x70, 0x70, 0x70)
USABLE_IN = 6.5  # letter, 1" margins


def _cell_bg(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexc)
    tcPr.append(shd)


def _fixed_layout(table):
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)


def _runs(paragraph, text, base_size=None):
    text = text.replace("~~", "").replace("\\*", "\u2217").replace("\\`", "`")
    text = re.sub(r"\\([\[\]()#_.-])", r"\1", text)
    for part in re.split(r"(\*\*[^*]+\*\*|(?<!\w)\*[^*\n]+\*(?!\w)|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"
            r.font.size = Pt((base_size or 10) - 1.5)
            continue
        else:
            r = paragraph.add_run(part)
        if base_size:
            r.font.size = Pt(base_size)


def _col_weights(header):
    n = len(header)
    w = [1.0] * n
    if header and header[0].strip() in ("#", "No", "No.", "Element", "Dimension", "Section"):
        w[0] = 0.42
    w[-1] = 1.7  # last column (verification / why / notes) tends to be long
    if n >= 5:
        w[-2] = 1.3
    tot = sum(w)
    return [x / tot for x in w]


def _emit_table(doc, rows):
    rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    header = rows[0] + [""] * (ncol - len(rows[0]))
    weights = _col_weights(header)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _fixed_layout(t)
    widths = [Inches(USABLE_IN * wt) for wt in weights]
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            cell.width = widths[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            _runs(p, row[ci] if ci < len(row) else "", base_size=8.5)
            if ri == 0:
                for rn in p.runs:
                    rn.bold = True; rn.font.color.rgb = RGBColor(255, 255, 255)
                _cell_bg(cell, "1F4E79")
    for ci in range(ncol):
        t.columns[ci].width = widths[ci]
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md_path, docx_path, title):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(0.9)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.05
    for s, sz in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)):
        stt = doc.styles[s]
        stt.font.color.rgb = ACCENT; stt.font.size = Pt(sz)
        stt.paragraph_format.space_before = Pt(10); stt.paragraph_format.space_after = Pt(4)

    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run(title); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = ACCENT
    doc.add_paragraph()

    i, in_fence, fence_buf, seen_h1, note_buf = 0, False, [], False, []

    def flush_notes():
        if note_buf:
            txt = re.sub(r"\*\*|\*|`", "", " ".join(note_buf))
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(txt); run.italic = True; run.font.color.rgb = NOTE
            note_buf.clear()

    while i < len(lines):
        raw = lines[i]
        if raw.strip().startswith("```"):
            if in_fence:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(fence_buf))
                run.font.name = "Consolas"; run.font.size = Pt(8.5)
                fence_buf = []; in_fence = False
            else:
                flush_notes(); in_fence = True
            i += 1; continue
        if in_fence:
            fence_buf.append(raw); i += 1; continue

        s = raw.strip()
        if s.startswith(">"):
            note_buf.append(re.sub(r"^>\s?", "", s)); i += 1; continue
        flush_notes()

        if not s:
            i += 1; continue
        if s == "---":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BBBBBB")):
                bot.set(qn(k), v)
            b.append(bot); pPr.append(b)
            i += 1; continue
        if s.startswith("# "):
            if not seen_h1:
                seen_h1 = True; i += 1; continue   # dedupe: title already added
            doc.add_heading(s[2:], level=1); i += 1; continue
        if s.startswith("## "):
            doc.add_heading(s[3:].replace("*", ""), level=2); i += 1; continue
        if s.startswith("### "):
            doc.add_heading(s[4:].replace("*", ""), level=3); i += 1; continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _emit_table(doc, rows); continue
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            _runs(p, re.sub(r"^[-*]\s+", "", s)); i += 1; continue
        if re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            _runs(p, re.sub(r"^\d+\.\s+", "", s)); i += 1; continue
        p = doc.add_paragraph(); _runs(p, s); i += 1

    flush_notes()
    doc.save(docx_path)
    print("wrote", docx_path)


if __name__ == "__main__":
    b = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "report")
    convert(b + r"\REPORT.md", b + r"\CT3201_Report.docx",
            "CT3201 Mini Project — Project Report")
    convert(b + r"\AI_INTERACTION_LOG.md", b + r"\CT3201_AI_Interaction_Log.docx",
            "CT3201 Mini Project — AI Interaction Log")
